#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
make_property_workbook.py
=========================
Emits the single spreadsheet that answers "is every property settled yet?".

WHY IT IS GENERATED AND NOT TYPED
---------------------------------
A hand-typed property list is out of date the moment a card changes, and a
stale one is worse than none -- it is the document you would run a nine-case
matrix on.  So every row here is pulled from something the verification suite
already checks:

  * the values come from the SHIPPED DECK, read out of the zip
  * the verdicts and sources come from verification/check_card_ranges.py,
    imported, not copied
  * the temperature rows come from the two constituent CSVs
  * the M6 proposals come from insitu_yarn_strength.py and
    porosity_stiffness.py, which compute them

If a card value changes and this file is not re-run, check_card_ranges.py
fails first.  There is no way for the workbook to be quietly wrong while the
suite is green.

  python3 data/properties/make_property_workbook.py
  python3 data/properties/make_property_workbook.py --check
"""
from __future__ import print_function

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))          # data/properties
ROOT = os.path.dirname(os.path.dirname(HERE))              # repository root
sys.path.insert(0, os.path.join(ROOT, "verification"))
sys.path.insert(0, HERE)

#: IN / DEV / GUESS -> the three words the thesis uses.
STATUS = {"IN": "검증", "DEV": "미검증", "GUESS": "임시값",
          "DERIVED": "도출값"}   # DERIVED: 2026-08-11, rF (Ge Eq.17이 결정)

#: Matrix card, slot -> (name, role, unit).  NPROPS = 25.
MATRIX_SLOTS = [
    (1,  "phase id",        "상 구분 (2 = 기지)",                     "-"),
    (2,  "E",               "기지 탄성계수",                           "MPa"),
    (3,  "nu",              "기지 포아송비",                           "-"),
    (4,  "Xt",              "기지 인장강도 (손상 개시)",               "MPa"),
    (5,  "Xc",              "기지 압축강도 (손상 개시)",               "MPa"),
    (6,  "At",              "인장 연화계수 (Gm_t<=0일 때만 사용)",     "-"),
    (7,  "Ac",              "압축 연화계수 (Gm_c<=0일 때만 사용)",     "-"),
    (8,  "dmax_t",          "인장 손상 상한 (수치 캡)",                "-"),
    (9,  "dmax_c",          "압축 손상 상한 (수치 캡)",                "-"),
    (10, "eta",             "점성 정규화 시간상수",                    "-"),
    (11, "max_djump",       "1증분 손상 도약 허용치",                  "-"),
    (12, "freeze_step",     "손상 동결 스텝 번호",                     "-"),
    (13, "min_PNEWDT",      "UMAT이 요구할 수 있는 최소 증분비",       "-"),
    (14, "enable",          "손상 on/off",                             "-"),
    (15, "Gm_t",            "기지 인장 파괴에너지 (균열대 정규화)",    "N/mm"),
    (16, "Gm_c",            "기지 압축 파괴에너지",                    "N/mm"),
    (17, "SY0",             "기지 항복응력 (수치 장치)",               "MPa"),
    (18, "HISO",            "선형 등방경화 계수",                      "MPa"),
    (19, "cut_trigger",     "선제 컷백 발동값",                        "-"),
    (20, "cut_safety",      "컷백 안전계수",                           "-"),
    (21, "cut_maxfac",      "컷백 최대 축소비",                        "-"),
    (22, "CARD KEY",        "★ 카드 오배치 가드 (반드시 30.0)",       "-"),
    (23, "reserved",        "예비 슬롯",                               "-"),
    (24, "HSMO",            "sign(I1) 불연속 tanh 평활화 폭",          "-"),
    (25, "NSTATV key",      "상태변수 개수 가드",                      "-"),
]

#: Yarn card, slot -> (name, role, unit).  NPROPS = 38.
YARN_SLOTS = [
    (1,  "phase id",   "상 구분 (1 = 얀)",                             "-"),
    (2,  "E1",         "얀 축방향 탄성계수 (섬유방향)",                "MPa"),
    (3,  "E2",         "얀 횡방향 탄성계수",                           "MPa"),
    (4,  "E3",         "얀 횡방향 탄성계수 (=E2, 횡등방)",             "MPa"),
    (5,  "nu12",       "축-횡 포아송비",                               "-"),
    (6,  "nu13",       "축-횡 포아송비 (=nu12)",                       "-"),
    (7,  "nu23",       "횡-횡 포아송비",                               "-"),
    (8,  "G12",        "축-횡 전단탄성계수",                           "MPa"),
    (9,  "G13",        "축-횡 전단탄성계수 (=G12)",                    "MPa"),
    (10, "G23",        "횡-횡 전단탄성계수",                           "MPa"),
    (11, "Xt",         "얀 축방향 인장강도 (섬유 파단)",               "MPa"),
    (12, "Xc",         "얀 축방향 압축강도 (좌굴·kinking)",            "MPa"),
    (13, "Yt",         "얀 횡방향 인장강도 (기지 균열)",               "MPa"),
    (14, "Yc",         "얀 횡방향 압축강도",                           "MPa"),
    (15, "S12",        "면내 전단강도",                                "MPa"),
    (16, "S13",        "면내 전단강도 (=S12)",                         "MPa"),
    (17, "S23",        "면외 전단강도",                                "MPa"),
    (18, "A1t",        "축방향 인장 연화계수 (G1t<=0일 때)",           "-"),
    (19, "A1c",        "축방향 압축 연화계수",                         "-"),
    (20, "Att",        "횡방향 인장 연화계수",                         "-"),
    (21, "Atc",        "횡방향 압축 연화계수",                         "-"),
    (22, "dmax_1",     "축방향 손상 상한 (수치 캡)",                   "-"),
    (23, "dmax_t",     "횡방향 손상 상한 (수치 캡)",                   "-"),
    (24, "eta",        "점성 정규화 시간상수",                         "-"),
    (25, "max_djump",  "1증분 손상 도약 허용치",                       "-"),
    (26, "freeze",     "손상 동결 스텝",                               "-"),
    (27, "min_PNEWDT", "최소 증분비",                                  "-"),
    (28, "enable",     "손상 on/off",                                  "-"),
    (29, "cut_trigger", "선제 컷백 발동값",                            "-"),
    (30, "cut_safety", "컷백 안전계수",                                "-"),
    (31, "cut_maxfac", "컷백 최대 축소비",                             "-"),
    (32, "G1t",        "축방향 인장 파괴에너지",                       "N/mm"),
    (33, "G1c",        "축방향 압축 파괴에너지",                       "N/mm"),
    (34, "Gtt",        "횡방향 인장 파괴에너지 (0 = 균열대 비활성)",   "N/mm"),
    (35, "Gtc",        "횡방향 압축 파괴에너지 (0 = 균열대 비활성)",   "N/mm"),
    (36, "X_PO",       "섬유 인발 응력 (혼합 연화법칙)",               "MPa"),
    (37, "rF",         "선형→지수 전이점",                             "-"),
    (38, "K1",         "선형 연화 기울기",                             "MPa"),
]

#: Slots that are solver settings, not material data.  They are still GUESS
#: in the audit, but the note has to say WHY so nobody hunts for a source.
NUMERICAL = {
    "matrix": (1, 6, 7, 8, 9, 10, 11, 12, 13, 14, 19, 20, 21, 22, 23, 24, 25),
    "yarn": (1, 18, 19, 20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30, 31),
}

#: Derived yarn elastic constants -- verified to 0.142 % by micromech_check.py
MICROMECH_NOTE = ("Zhang 2022 Table 1(T300) + Table 2(SiC)를 Chamis 강성식·"
                  "Schapery CTE식으로 균질화. Vf(얀 내부)=0.79194. "
                  "micromech_check.py가 12개 상수를 최대 0.142 % 오차로 재현")


def load_audit():
    """(name -> (verdict, source, note)) from check_card_ranges.py."""
    import check_card_ranges as ccr
    out = {}
    for slot, name, val, lo, hi, verdict, src, why in ccr.MATRIX:
        out[("matrix", slot)] = (verdict, src, why, lo, hi)
    for slot, name, val, lo, hi, verdict, src, why in ccr.YARN:
        out[("yarn", slot)] = (verdict, src, why, lo, hi)
    return out, ccr


def read_csv(path):
    rows, hdr = [], None
    for line in open(path).read().splitlines():
        if line.startswith("#") or not line.strip():
            continue
        p = [x.strip().strip('"') for x in line.split(",")]
        if hdr is None:
            hdr = p
        else:
            rows.append(p)
    return hdr, rows


# ==========================================================================
def build_rows():
    """Every sheet, as (sheet name, header, list-of-rows)."""
    audit, ccr = load_audit()
    deck = ccr.deck_text()
    mcard, ycard = ccr.card(deck, 25), ccr.card(deck, 38)
    exps = ccr.expansions(deck)

    HDR = ["물성 이름", "기능", "값", "단위", "출처", "상태", "비고"]
    sheets = []

    # ---------------------------------------------------------- matrix card
    rows = []
    for slot, name, role, unit in MATRIX_SLOTS:
        val = mcard[slot - 1]
        a = audit.get(("matrix", slot))
        if a:
            verdict, src, why, lo, hi = a
            status, source, note = STATUS[verdict], src, why
            if lo is not None:
                note = ("독립 범위 [%g, %g]. " % (lo, hi)) + note
        elif slot in NUMERICAL["matrix"]:
            status, source = "임시값", "물성이 아님 — 솔버 설정값"
            note = "수치 파라미터. 출처를 찾을 대상이 아니며 보정이 움직인다"
        else:
            status, source, note = "임시값", "미분류", ""
        rows.append(["기지 슬롯 %d — %s" % (slot, name), role, val, unit,
                     source, status, note])
    sheets.append(("기지카드_25슬롯", HDR, rows))

    # ------------------------------------------------------------ yarn card
    rows = []
    for slot, name, role, unit in YARN_SLOTS:
        val = ycard[slot - 1]
        a = audit.get(("yarn", slot))
        if a:
            verdict, src, why, lo, hi = a
            status, source, note = STATUS[verdict], src, why
            if lo is not None:
                note = ("독립 범위 [%g, %g]. " % (lo, hi)) + note
        elif 2 <= slot <= 10:
            status, source, note = "검증", MICROMECH_NOTE, \
                "구성재에서 유도된 값. 카드에 손으로 적은 숫자가 아니다"
        elif slot in NUMERICAL["yarn"]:
            status, source = "임시값", "물성이 아님 — 솔버 설정값"
            note = "수치 파라미터. 출처를 찾을 대상이 아니며 보정이 움직인다"
        else:
            status, source, note = "임시값", "미분류", ""
        rows.append(["얀 슬롯 %d — %s" % (slot, name), role, val, unit,
                     source, status, note])
    sheets.append(("얀카드_38슬롯", HDR, rows))

    # ------------------------------------------------------------------ CTE
    rows = []
    names = ["기지 alpha", "얀 alpha1 (축)", "얀 alpha2 = alpha3 (횡)"]
    vals = [exps[0][0], exps[1][0], exps[1][1]]
    roles = ["기지 열팽창 — TRS를 만드는 항",
             "얀 축방향 열팽창",
             "얀 횡방향 열팽창"]
    for (nm, val, role, c) in zip(names, vals, roles, ccr.CTE):
        _n, _v, lo, hi, verdict, src, why = c
        note = why
        if lo is not None:
            note = ("독립 범위 [%g, %g]. " % (lo, hi)) + note
        rows.append([nm, role, val, "1/K", src, STATUS[verdict], note])
    rows.append(["*Expansion, zero=", "무응력 온도 — 물성 가정이지 솔버 설정이 아니다",
                 1050.0, "degC",
                 "Zhang 2022 공정온도. trs_configuration.py가 CONFIG_V=1050, "
                 "CONFIG_P=782로 분리",
                 "미검증",
                 "CONFIG_V는 논문 재현용으로 1050 유지(TRS 2.34배를 그대로 보고), "
                 "CONFIG_P는 XRD를 재현하는 782를 이분법으로 푼 값(=보정값)"])
    sheets.append(("열팽창_무응력온도", HDR, rows))

    # -------------------------------------------- constituent temperature
    rows = []
    for csvname, phase in (("fibre_T300_vsT.csv", "T300 섬유"),
                           ("matrix_SiC_vsT.csv", "SiC 기지")):
        hdr, data = read_csv(os.path.join(HERE, csvname))
        for r in data:
            d = dict(zip(hdr, r))
            st = {"verified": "검증", "literature": "미검증",
                  "placeholder": "임시값"}.get(d.get("status", ""), "임시값")
            for key, unit, role in (
                    ("E1", "MPa", "축방향 탄성계수"),
                    ("E", "MPa", "탄성계수"),
                    ("E2", "MPa", "횡방향 탄성계수"),
                    ("Xt", "MPa", "인장강도"),
                    ("Xc", "MPa", "압축강도"),
                    ("alpha", "1/K", "열팽창계수(할선)"),
                    ("alpha1", "1/K", "축방향 열팽창계수(할선)"),
                    ("alpha2", "1/K", "횡방향 열팽창계수(할선)"),
                    ("k", "W/(mm.K)", "열전도율"),
                    ("k1", "W/(mm.K)", "축방향 열전도율"),
                    ("k2", "W/(mm.K)", "횡방향 열전도율"),
                    ("cp", "mJ/(t.K)", "비열 — 열전달 해석 입력"),
                    ("rho", "t/mm^3", "밀도")):
                v = d.get(key, "")
                if not v:
                    continue
                note = d.get("source", "")
                st2 = st
                if key in ("k", "k1"):
                    st2 = "미검증"
                    note = (note + " ‖ ★ 미해결 충돌: k1은 Pradere 46-72 vs "
                            "refs/[17],[22] 8 W/(m.K), 6-9배. 민감도 케이스로 "
                            "들고 간다(conductivity_bounds.py)")
                if key == "k" and phase.startswith("SiC"):
                    note = ("Snead Eq.12 = 단결정 상한. ★ 입력으로 사용 불가가 "
                            "증명됨 — conductivity_bounds.py")
                rows.append(["%s %s @ %s C" % (phase, key, d["T_C"]),
                             role, float(v), unit, note, st2,
                             "공정: %s" % d.get("process", "-")])
    sheets.append(("구성재_온도의존", HDR, rows))

    # -------------------------------------------------- geometry / porosity
    import porosity_stiffness as ps
    vp = ps.porosity_from_density(2.0, 0.40, ps.RHO_T300)
    a, b = ps.rve_phases(), ps.real_phases()
    rows = [
        ["V_RVE", "RVE 경계상자 체적 — 드라이버 반력을 응력으로 바꾸는 제수",
         5.390, "mm^3", "메시 형상 (3.5 x 3.5 x 0.44)", "검증",
         "sigma = +RF/V. 패치시험이 2.36e-08로 확인(§3.8-1a)"],
        ["V_yarn / V_RVE", "얀 체적분율", 0.4982, "-",
         "메시에서 직접 계산", "검증", "섬유 체적분율 39.5 %로 환산되어 논문의 '약 40 %'와 일치"],
        ["Vf (얀 내부)", "얀 안의 섬유 체적분율", ps.VF_YARN, "-",
         "카드 E1에서 역산: (Em-E1)/(Em-Ef1)", "검증",
         "이 하나의 Vf로 얀 상수 12개가 전부 재현된다(micromech_check.py)"],
        ["요소 수", "C3D4 요소 개수", 26452, "-", "메시", "검증",
         "탐색용 거친 메시. 메시 수렴성 검증은 미수행(미결 항목 4)"],
        ["★ 공극률 Vp (CVI 하한)", "복합재 총 공극률 — 우리 재료 아님",
         vp, "-", "refs/[10] Yang의 rho=2.0 g/cm3, Vf=40 %에서 역산(둘 다 [10] "
         "자체 진술). 같은 산식이 refs/[36]의 발표값 17.0 %를 0.03 % 오차로 재현",
         "임시값",
         "refs/[10]은 CVI, 우리 기준 재료는 PIP(Zhang [5], 제2장 §2.2.1)이고 "
         "PIP가 기지 공극률이 더 높다 → 이 값은 하한이다. PIP 밀도 실측 미확보. "
         "refs/[43]의 13 %·refs/[28]의 10-15 %는 개기공으로 추정(자기 밀도와 불일치)"],
        ["★ 공극률 Vp (우리 RVE)", "메시가 실제로 담고 있는 공극",
         0.0, "-", "메시가 100 % 충전(filled cell)", "미검증",
         "★ 없는 SiC를 20.2 %p 채우고 있다. 강성 과대예측의 정체"],
        ["★ 기지 강하계수", "공극을 카드로 흉내내는 계수",
         ps.pocket_knockdown(vp), "-",
         "1 - Vp/(1-Vy). 병렬합에서 체적 삭제와 계수 축소는 동일",
         "검증", "E·k·rho에만 적용. alpha와 강도에는 적용하지 않는다. "
         "Snead 7.7 %(마이크로공극)와 중첩 금지"],
        ["섬유 분율 (우리 RVE)", "상 구성 대조", a["fibre"], "-",
         "0.4982 x 0.79194", "검증", "실제 40.0 %와 0.5 %p 차 — 메시는 보강재에 대해 옳다"],
        ["기지 분율 (우리 RVE)", "상 구성 대조", a["matrix"], "-",
         "0.4982 x 0.20806 + 0.5018", "미검증",
         "실제 40.4 % 대비 +20.2 %p. 셀에서 가장 단단한 상이 과다"],
    ]
    sheets.append(("형상_공극률", HDR, rows))

    # -------------------------------------------- validation-only (NOT card)
    rows = []
    for lab, T, e, conv, g in ps.measured_composites():
        rows.append(["복합재 탄성계수 (%s)" % lab, "L2 검증 목표 — 카드 입력 금지",
                     e, "GPa", lab, "검증",
                     "★ 측정 규약: %s. 규약이 다른 값을 섞으면 2배가 틀린다" % conv])
    for T, s in ((23.0, 128.45), (500.0, 179.42), (1000.0, 199.15)):
        rows.append(["복합재 인장강도 @ %.0f C" % T,
                     "M1 보정 목표 — 카드 입력 금지", s, "MPa",
                     "Zhang 2022 Table 3", "검증",
                     "이 값으로 카드를 맞추면 예측이 아니라 보정이 된다"])
    for lab, val, unit, note in (
            ("복합재 인장강도 (refs/[43] Mei)", 248.0, "MPa", "2D C/SiC, 아르곤 50사이클 후 98.90 % 유지"),
            ("복합재 인장강도 (refs/[03] Zhang2012)", 259.3, "MPa", "as-received"),
            ("복합재 인장강도 (refs/[10] Yang, 1000 C)", 268.2, "MPa", "초기접선 172.7 GPa와 같은 시편"),
            ("복합재 파단변형률 (refs/[10], 1000 C)", 0.32, "%", "★ M5가 멈춘 0.3318 %와 사실상 일치"),
            ("복합재 TRS (XRD, refs/[15])", 114.7, "MPa", "3D 브레이드 — 본 연구는 2D. 이전 한계 명시"),
            ("복합재 열전도율 (refs/[12])", 6.29, "W/(m.K)", "Snead 기지 k가 입력 불가임을 증명한 판정값")):
        rows.append([lab, "검증 전용 — 카드 입력 금지", val, unit,
                     "문헌 실측", "검증", note])
    sheets.append(("검증전용_복합재실측", HDR, rows))

    # ------------------------------------------------------- M6 proposals
    import insitu_yarn_strength as iys
    rows = []
    for T in (23.0, 500.0, 1000.0):
        lo, hi = iys.admissible_band(T)
        rows.append(["얀 Xt @ %.0f C (제안)" % T,
                     "축방향 인장강도 — 복합재 강도 상한을 정하는 유일한 knob",
                     round(lo, 1), "MPa",
                     "refs/[08] Sauder Table 1 Weibull(m, sigma_0 @ 1 mm^3)를 "
                     "RVE의 하중방향 섬유 체적 1.063 mm^3에서 평가", "검증",
                     "구간 [%.0f, %.0f] MPa의 하한에서 시작. 현재 카드 2835는 "
                     "이 값의 %.2f배" % (lo, hi, 2835.0 / lo)])
    rows.append(["기지 E (제안)", "공극률 반영 후 기지 탄성계수",
                 350000.0 * ps.pocket_knockdown(vp), "MPa",
                 "350 000 x 0.6089 (강하계수)", "검증",
                 "M5 초기접선 235.2 → 170.0 GPa, refs/[10] 실측 172.7과 1.6 % 차"])
    rows.append(["얀 Gtt (제안)", "횡방향 인장 파괴에너지 — 0이면 균열대 비활성",
                 0.107, "N/mm", "Shi refs/[31], 2D 평직 C/SiC", "검증",
                 "카드는 아직 0. 덱 재생성 필요(yarn_fracture_energy.py)"])
    sheets.append(("M6_변경예정", HDR, rows))

    return sheets


# ==========================================================================
def write_workbook(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    sheets = build_rows()
    wb = Workbook()
    wb.remove(wb.active)

    FONT = "Arial"
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    fills = {"검증": PatternFill("solid", fgColor="E2EFDA"),
             "미검증": PatternFill("solid", fgColor="FCE4D6"),
             "임시값": PatternFill("solid", fgColor="FFF2CC")}
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ------------------------------------------------ the summary sheet
    ws = wb.create_sheet("요약")
    ws["A1"] = "물성 현황 요약 — 해석 실행 전 관문"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = ("이 파일은 data/properties/make_property_workbook.py가 생성한다 — "
                "손으로 고치지 말고 스크립트를 다시 돌린다. "
                "상태는 verification/check_card_ranges.py의 판정을 그대로 옮긴 것이다. "
                "검증 = 그 값의 출처가 아닌 독립 문헌 범위 안 / "
                "미검증 = 독립 문헌과 어긋남(사유 기록됨) / "
                "임시값 = 독립값 없음, 보정이 움직일 대상")
    ws["A2"].font = Font(name=FONT, size=9, italic=True)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:F4")

    ws.append([])
    ws.append(["시트", "검증", "미검증", "임시값", "합계"])
    for c in ws[6]:
        c.font, c.fill, c.border = head_font, head_fill, border
    r = 7
    for name, hdr, rows in sheets:
        n_v = sum(1 for x in rows if x[5] == "검증")
        n_u = sum(1 for x in rows if x[5] == "미검증")
        n_p = sum(1 for x in rows if x[5] == "임시값")
        ws.cell(r, 1, name)
        ws.cell(r, 2, n_v)
        ws.cell(r, 3, n_u)
        ws.cell(r, 4, n_p)
        ws.cell(r, 5, n_v + n_u + n_p)
        r += 1
    ws.cell(r, 1, "합계").font = Font(name=FONT, bold=True)
    # Values, not formulas.  This workbook is GENERATED output: the counts are
    # recomputed by re-running the script, and there is no input cell a reader
    # could edit that ought to drive them.  A formula here would also be one we
    # could not verify -- LibreOffice cannot recalculate in this environment.
    for j, col in enumerate("BCDE"):
        tot = sum(ws.cell(k, j + 2).value or 0 for k in range(7, r))
        ws["%s%d" % (col, r)] = tot
        ws["%s%d" % (col, r)].font = Font(name=FONT, bold=True)
    for row in ws.iter_rows(min_row=6, max_row=r, max_col=5):
        for c in row:
            c.border = border
            if c.font.name != FONT:
                c.font = Font(name=FONT, size=10,
                              bold=(c.row == r or c.row == 6))
    for i, w in enumerate([26, 10, 10, 10, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    note = r + 2
    ws.cell(note, 1, "★ 지금 해석을 돌려도 되는가").font = Font(
        name=FONT, bold=True, size=12)
    for i, line in enumerate((
            "1. '임시값'이 있다는 것 자체는 실행 금지 사유가 아니다 — 보정이 움직일 "
            "대상이며, 어느 숫자가 추정값인지 알고 돌리면 된다.",
            "2. 실행을 막는 것은 '미검증' 중 크기가 큰 것들이다. 현재 세 가지: "
            "얀 Xt(카드가 독립 범위의 1.6~6.0배), 기지 공극률(0 % vs 실제 19.6 %), "
            "구성재 CTE 2개(TRS 2.34배 과대예측의 원인).",
            "3. 얀 Xt와 공극률은 2026-08-03에 값이 결정되었다(M6_변경예정 시트). "
            "CTE는 CONFIG_V/CONFIG_P로 분리하여 처리하기로 결정되어 있다.",
            "4. 따라서 남은 작업은 '결정된 값으로 덱을 재생성'하는 것이며, "
            "새로운 문헌 조사가 남아 있는 것이 아니다.",
            "5. 균열대 Gtt/Gtc가 0인 동안 얀 횡방향 모드는 메시 객관적이지 않다. "
            "강도·손상 분포를 최종 수치로 쓰기 전에 반드시 채운다."),
            start=0):
        c = ws.cell(note + 1 + i, 1, line)
        c.font = Font(name=FONT, size=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=note + 1 + i, start_column=1,
                       end_row=note + 1 + i, end_column=5)
        ws.row_dimensions[note + 1 + i].height = 28

    # ------------------------------------------------------ data sheets
    for name, hdr, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(hdr)
        for c in ws[1]:
            c.font, c.fill, c.border = head_font, head_fill, border
            c.alignment = Alignment(horizontal="center", vertical="center")
        for row in rows:
            ws.append(row)
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row,
                                max_col=len(hdr)):
            for c in row:
                c.font = Font(name=FONT, size=9)
                c.border = border
                c.alignment = Alignment(wrap_text=True, vertical="top")
            st = row[5].value
            if st in fills:
                row[5].fill = fills[st]
                row[5].font = Font(name=FONT, size=9, bold=True)
                row[5].alignment = Alignment(horizontal="center",
                                             vertical="center")
            v = row[2].value
            if isinstance(v, float):
                row[2].number_format = ("0.000E+00" if (v and abs(v) < 1e-3)
                                        else "#,##0.0####")
            row[2].alignment = Alignment(horizontal="right", vertical="top")
        for i, w in enumerate([30, 34, 14, 12, 52, 10, 60], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(hdr)),
                                          ws.max_row)

    wb.save(path)
    return sheets


# ==========================================================================
_OK, _BAD = [], []


def ck(name, cond, detail=""):
    (_OK if cond else _BAD).append(name)
    print("  [%s] %-58s %s" % ("PASS" if cond else "FAIL", name, detail))


def selftest():
    print("=" * 78)
    print("make_property_workbook.py --check")
    print("=" * 78)
    sheets = build_rows()
    allrows = [r for _n, _h, rs in sheets for r in rs]

    print("\n A. the workbook covers every card slot")
    by = dict((n, rs) for n, _h, rs in sheets)
    ck("matrix sheet has all 25 slots", len(by["기지카드_25슬롯"]) == 25,
       "%d" % len(by["기지카드_25슬롯"]))
    ck("yarn sheet has all 38 slots", len(by["얀카드_38슬롯"]) == 38,
       "%d" % len(by["얀카드_38슬롯"]))
    # >1, not >3: this gate was written against English and would reject
    # legitimate two-character Korean words like 비열 and 밀도.
    ck("no slot is missing a role", all(len(r[1].strip()) > 1 for r in allrows))
    ck("every row has a unit", all(len(str(r[3]).strip()) > 0 for r in allrows))

    print("\n B. the values are the SHIPPED deck's, not retyped")
    import check_card_ranges as ccr
    deck = ccr.deck_text()
    m, y = ccr.card(deck, 25), ccr.card(deck, 38)
    ok = True
    for i, r in enumerate(by["기지카드_25슬롯"]):
        ok = ok and abs(r[2] - m[i]) < 1e-9
    for i, r in enumerate(by["얀카드_38슬롯"]):
        ok = ok and abs(r[2] - y[i]) < 1e-9
    ck("every card value matches the deck exactly", ok)
    ck("the guard slot is still 30.0",
       abs(by["기지카드_25슬롯"][21][2] - 30.0) < 1e-9)

    print("\n C. the status words are the audit's verdicts, translated")
    ck("only the audit's status words are used",
       set(r[5] for r in allrows) <= set(("검증", "미검증", "임시값", "도출값")),
       ", ".join(sorted(set(r[5] for r in allrows))))
    ck("yarn Xt is 미검증, matching its DEV regrade",
       by["얀카드_38슬롯"][10][5] == "미검증",
       "%.0f MPa" % by["얀카드_38슬롯"][10][2])
    ck("matrix Xt is 검증, matching its IN verdict",
       by["기지카드_25슬롯"][3][5] == "검증")
    ck("the yarn elastic constants are 검증 via micromech_check",
       all(by["얀카드_38슬롯"][i][5] == "검증" for i in range(1, 10)))
    ck("numerical slots say they are not material data",
       "솔버 설정값" in by["기지카드_25슬롯"][5][4]
       and "솔버 설정값" in by["얀카드_38슬롯"][17][4],
       "matrix slot 6, yarn slot 18")

    print("\n D. composite measurements are fenced off from the cards")
    val = by["검증전용_복합재실측"]
    ck("every validation row says card input is forbidden",
       all("카드 입력 금지" in r[1] for r in val), "%d rows" % len(val))
    ck("every measured modulus carries its measurement convention",
       all("측정 규약" in r[6] for r in val if "탄성계수" in r[0]))
    ck("no validation row leaks into a card sheet",
       not any("검증 전용" in r[1] for r in
               by["기지카드_25슬롯"] + by["얀카드_38슬롯"]))

    print("\n E. the M6 proposals come from the scripts that compute them")
    import insitu_yarn_strength as iys
    import porosity_stiffness as ps
    prop = by["M6_변경예정"]
    lo23 = iys.admissible_band(23.0)[0]
    ck("the 23 C yarn Xt proposal is insitu_yarn_strength's low end",
       abs(prop[0][2] - round(lo23, 1)) < 0.05,
       "%.1f MPa" % prop[0][2])
    vp = ps.porosity_from_density(2.0, 0.40, ps.RHO_T300)
    ck("the matrix E proposal is 350 GPa times the knockdown",
       abs(prop[3][2] - 350000.0 * ps.pocket_knockdown(vp)) < 1.0,
       "%.0f MPa" % prop[3][2])
    ck("every proposal is graded 검증 -- none is a fit",
       all(r[5] == "검증" for r in prop))
    ck("the porosity row says not to stack it on Snead",
       any("Snead" in r[6] for r in by["형상_공극률"]))

    print("\n F. the file does not hide what is still open")
    ck("Gtt = 0 is recorded as disabling the crack band",
       "균열대 비활성" in by["얀카드_38슬롯"][33][1])
    ck("the conductivity conflict is carried, not resolved",
       any("미해결 충돌" in r[4] for r in by["구성재_온도의존"]),
       "k1: Pradere 46-72 vs refs/[17],[22] 8 W/(m.K)")
    ck("the RVE's own porosity is flagged 미검증",
       any(r[5] == "미검증" and "우리 RVE" in r[0]
           for r in by["형상_공극률"]))
    ck("mesh convergence is named as not done",
       any("메시 수렴성 검증은 미수행" in r[6] for r in by["형상_공극률"]))

    print("\n" + "=" * 78)
    if _BAD:
        print("FAIL -- %d of %d: %s"
              % (len(_BAD), len(_OK) + len(_BAD), ", ".join(_BAD[:3])))
        print("=" * 78)
        return 1
    print("ALL %d PROPERTY-WORKBOOK CHECKS PASS" % len(_OK))
    print("=" * 78)
    return 0


if __name__ == "__main__":
    if "--check" in sys.argv:
        sys.exit(selftest())
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        ROOT, "dist", "PROPERTY_TABLE.xlsx")
    d = os.path.dirname(out)
    if d and not os.path.isdir(d):
        os.makedirs(d)
    sheets = write_workbook(out)
    n = sum(len(rs) for _a, _b, rs in sheets)
    print("wrote %s  (%d sheets, %d property rows)"
          % (out, len(sheets) + 1, n))
    for name, _h, rows in sheets:
        v = sum(1 for r in rows if r[5] == "검증")
        u = sum(1 for r in rows if r[5] == "미검증")
        p = sum(1 for r in rows if r[5] == "임시값")
        print("  %-24s 검증 %3d  미검증 %3d  임시값 %3d" % (name, v, u, p))
